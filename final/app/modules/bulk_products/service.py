import csv
import io
from datetime import datetime
from openpyxl import Workbook, load_workbook
from fastapi import HTTPException

from app.modules.products.service import create_product_service, list_products_service

FIELDS = [
    "sku","name","supplier_id","purchase_price","international_shipping",
    "domestic_shipping","selling_price","fee_rate","vat_rate",
    "main_image_url","detail_image_url"
]
REQUIRED = {"name"}
NUMERIC = {"purchase_price","international_shipping","domestic_shipping","selling_price","fee_rate","vat_rate"}

def _clean(row: dict) -> dict:
    data = {}
    for key in FIELDS:
        value = row.get(key)
        if isinstance(value, str):
            value = value.strip()
        if value in ("", None):
            if key in NUMERIC:
                if key == "vat_rate":
                    value = 0.1
                elif key in {"purchase_price","international_shipping","domestic_shipping","selling_price","fee_rate"}:
                    value = 0
                else:
                    continue
            else:
                value = None
        data[key] = value
    return data

def parse_csv(content: bytes) -> list[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try: text = content.decode("cp949")
        except UnicodeDecodeError: raise HTTPException(status_code=422, detail="CSV encoding must be UTF-8 or CP949")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(status_code=422, detail="CSV header is required")
    missing = REQUIRED - set(reader.fieldnames)
    if missing:
        raise HTTPException(status_code=422, detail=f"missing required columns: {', '.join(sorted(missing))}")
    return [_clean(row) for row in reader]

def parse_xlsx(content: bytes) -> list[dict]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="invalid Excel file") from exc
    ws = wb.active
    values = ws.iter_rows(values_only=True)
    try: headers = [str(x).strip() if x is not None else "" for x in next(values)]
    except StopIteration: raise HTTPException(status_code=422, detail="Excel file is empty")
    missing = REQUIRED - set(headers)
    if missing:
        raise HTTPException(status_code=422, detail=f"missing required columns: {', '.join(sorted(missing))}")
    rows=[]
    for vals in values:
        if not any(v not in (None,"") for v in vals): continue
        rows.append(_clean(dict(zip(headers, vals))))
    return rows

async def import_products_service(user_id: str, filename: str, content: bytes):
    if not content: raise HTTPException(status_code=422, detail="file is empty")
    if len(content) > 10 * 1024 * 1024: raise HTTPException(status_code=413, detail="file must be 10 MB or smaller")
    if len(filename or "") > 255: raise HTTPException(status_code=422, detail="filename is too long")
    lower=(filename or "").lower()
    if lower.endswith(".csv"): rows=parse_csv(content)
    elif lower.endswith(".xlsx"): rows=parse_xlsx(content)
    else: raise HTTPException(status_code=422, detail="only .csv and .xlsx are supported")
    if len(rows) > 5000: raise HTTPException(status_code=422, detail="maximum 5000 rows per upload")

    successes=[]; errors=[]
    for index,row in enumerate(rows,start=2):
        try:
            product=await create_product_service(user_id,row)
            successes.append({"row":index,"product_id":product.get("id"),"sku":product.get("sku"),"name":product.get("name"),"status":product.get("status")})
        except HTTPException as exc:
            errors.append({"row":index,"name":row.get("name"),"sku":row.get("sku"),"error":str(exc.detail)})
        except Exception as exc:
            errors.append({"row":index,"name":row.get("name"),"sku":row.get("sku"),"error":"unexpected row error"})
    return {"total_rows":len(rows),"success_count":len(successes),"error_count":len(errors),"successes":successes,"errors":errors}

async def export_rows(user_id: str):
    products=await list_products_service(user_id)
    columns=FIELDS+["total_cost","fee_amount","vat_amount","net_profit","margin_rate","status","status_reason"]
    rows=[{key:p.get(key) for key in columns} for p in products]
    return columns,rows

async def export_csv_service(user_id: str) -> bytes:
    columns,rows=await export_rows(user_id)
    out=io.StringIO(); w=csv.DictWriter(out,fieldnames=columns); w.writeheader(); w.writerows(rows)
    return ("\ufeff"+out.getvalue()).encode("utf-8")

async def export_xlsx_service(user_id: str) -> bytes:
    columns,rows=await export_rows(user_id)
    wb=Workbook(); ws=wb.active; ws.title="products"; ws.append(columns)
    for row in rows: ws.append([row.get(c) for c in columns])
    out=io.BytesIO(); wb.save(out); return out.getvalue()

def template_csv() -> bytes:
    out=io.StringIO(); w=csv.DictWriter(out,fieldnames=FIELDS); w.writeheader()
    w.writerow({"sku":"SAMPLE-001","name":"샘플상품","purchase_price":10000,"international_shipping":1000,"domestic_shipping":3000,"selling_price":30000,"fee_rate":0.1,"vat_rate":0.1})
    return ("\ufeff"+out.getvalue()).encode("utf-8")

def template_xlsx() -> bytes:
    wb=Workbook(); ws=wb.active; ws.title="products"; ws.append(FIELDS)
    ws.append(["SAMPLE-001","샘플상품",None,10000,1000,3000,30000,0.1,0.1,None,None])
    out=io.BytesIO(); wb.save(out); return out.getvalue()
